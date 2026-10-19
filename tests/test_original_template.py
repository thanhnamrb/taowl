"""Regression checks for the original branded LingualXplore template.

The app's UI is top-level Streamlit code, so load only the model/renderer layer
without starting a Streamlit session.
"""
import ast
import io
import sys
import types
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
source = (ROOT / "app.py").read_text(encoding="utf-8")
module = ast.parse(source)
cut = source.index("# ===================== STREAMLIT UI =====================")
test_module = types.ModuleType("taowl_render_test")
sys.modules[test_module.__name__] = test_module
namespace = test_module.__dict__
exec(compile(source[:cut], str(ROOT / "app.py"), "exec"), namespace)

from docx import Document
from docx.oxml.ns import qn

class OriginalTemplateTests(unittest.TestCase):
    def setUp(self):
        self.template = ROOT / "templates" / "lx_vocab_template.docx"
        self.assertTrue(self.template.is_file())

    def sample(self, title="SPECIAL DAYS"):
        return namespace["parse_vocab_csv"](
            "No.,Word,Type,Pronunciation,Meaning\n"
            "1,impress,v,/ɪmˈpres/,gây ấn tượng\n"
            ",impression,n,/ɪmˈpreʃn/,ấn tượng\n"
            "2,protect,v,/prəˈtekt/,bảo vệ\n",
            unit="5.1", heading_unit="5", title=title,
        )

    def test_original_artwork_preserved(self):
        data = namespace["render_vocab_docx"](self.sample(), template_path=self.template).getvalue()
        original = Document(self.template)
        produced = Document(io.BytesIO(data))
        self.assertEqual(len(original.inline_shapes), len(produced.inline_shapes))
        self.assertEqual(len(original.sections), len(produced.sections))
        self.assertGreaterEqual(len(produced.tables), 2)
        self.assertIn("SPECIAL DAYS", produced.tables[0].rows[1].cells[2].text)

    def test_long_title_and_page_field(self):
        title = "A VERY LONG AND IMPORTANT COLLECTION OF ENGLISH WORD FAMILIES FOR ADVANCED CLASSROOM PRACTICE"
        data = namespace["render_vocab_docx"](self.sample(title), template_path=self.template).getvalue()
        produced = Document(io.BytesIO(data))
        self.assertIn(title, produced.tables[0].rows[1].cells[2].text)
        footer_xml = produced.sections[0].footer._element.xml
        self.assertIn(" PAGE ", footer_xml)
        self.assertIn("gradFill", footer_xml)
        self.assertIn("updateFields", produced.settings.element.xml)

    def test_round_trip_title_metadata(self):
        data = namespace["render_vocab_docx"](self.sample(), template_path=self.template).getvalue()
        meta, parsed_csv, notes = namespace["extract_old_vocab_docx"](data)
        self.assertEqual(meta["unit"], "5.1")
        self.assertEqual(meta["heading_unit"], "5")
        self.assertEqual(meta["title"], "SPECIAL DAYS")
        self.assertEqual(len(self.sample().families), len(namespace["parse_vocab_csv"](parsed_csv, unit=meta["unit"], title=meta["title"]).families))

    def test_csv_malformed_rows_are_not_silently_discarded(self):
        parser = namespace["parse_vocab_csv"]
        with self.assertRaises(ValueError):
            parser("1,word,n,ipa,meaning,EXTRA", unit="1", title="TEST")
        with self.assertRaises(ValueError):
            parser("1,,n,ipa,meaning", unit="1", title="TEST")

if __name__ == "__main__":
    unittest.main()
