from __future__ import annotations

import csv
import io
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


TWIPS_PER_INCH = 1440
EMU_PER_TWIP = 635


@dataclass(slots=True)
class VocabWord:
    word: str
    word_type: str = ""
    pronunciation: str = ""
    meaning: str = ""


@dataclass(slots=True)
class VocabFamily:
    number: str
    words: list[VocabWord] = field(default_factory=list)


@dataclass(slots=True)
class VocabDocument:
    unit: str
    title: str
    document_type: str = "VOCAB BUILDER"
    section_label: str = "UNIT"
    heading_unit: str = ""
    families: list[VocabFamily] = field(default_factory=list)

    @property
    def unit_badge(self) -> str:
        return (self.unit or "").strip() or "0"

    @property
    def display_heading_unit(self) -> str:
        if (self.heading_unit or "").strip():
            return self.heading_unit.strip()
        return (self.unit or "").split(".", 1)[0].strip() or self.unit

    @property
    def word_count(self) -> int:
        return sum(len(f.words) for f in self.families)

    @property
    def family_count(self) -> int:
        return len(self.families)


@dataclass(slots=True)
class TextStyle:
    font: str = "Arial"
    size: float = 10.5
    color: str = "000000"
    bold: bool = False
    italic: bool = False


@dataclass(slots=True)
class Theme:
    accent: str = "BF4E14"
    navy: str = "0F4761"
    pale: str = "E7A07F"
    blue: str = "7FA9BC"
    body: TextStyle = field(default_factory=lambda: TextStyle("Times New Roman", 12, "000000"))
    header: TextStyle = field(default_factory=lambda: TextStyle("Arial", 10.5, "FFFFFF", True))
    no_text: TextStyle = field(default_factory=lambda: TextStyle("Arial", 10, "FFFFFF", True))
    title: TextStyle = field(default_factory=lambda: TextStyle("Montserrat", 27, "000000", True))
    document_type: TextStyle = field(default_factory=lambda: TextStyle("Arial", 18, "BF4E14"))
    badge: TextStyle = field(default_factory=lambda: TextStyle("Arial", 44, "FFFFFF", True))
    footer: TextStyle = field(default_factory=lambda: TextStyle("Arial", 9.5, "0F4761"))
    page: TextStyle = field(default_factory=lambda: TextStyle("Arial", 10, "FFFFFF", True))
    phone: str = "0345286842"
    email: str = "email@yourcenter.com"


def _hex(value: str, fallback: str = "000000") -> str:
    v = (value or "").strip().lstrip("#").upper()
    return v if len(v) == 6 and all(ch in "0123456789ABCDEF" for ch in v) else fallback


def _apply_run_style(run, style: TextStyle) -> None:
    run.font.name = style.font
    run.font.size = Pt(float(style.size))
    run.font.bold = bool(style.bold)
    run.font.italic = bool(style.italic)
    run.font.color.rgb = RGBColor.from_string(_hex(style.color))
    rfonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), style.font)


def rows_to_document(
    rows: list[dict],
    *,
    unit: str,
    title: str,
    heading_unit: str = "",
    section_label: str = "UNIT",
    document_type: str = "VOCAB BUILDER",
) -> VocabDocument:
    doc = VocabDocument(
        unit=(unit or "").strip(),
        title=(title or "").strip(),
        heading_unit=(heading_unit or "").strip(),
        section_label=(section_label or "UNIT").strip().upper(),
        document_type=(document_type or "VOCAB BUILDER").strip().upper(),
    )
    current: VocabFamily | None = None
    for row in rows:
        no = str(row.get("No.", "") or "").strip()
        word = str(row.get("Word", "") or "").strip()
        if not word:
            continue
        if no or current is None:
            current = VocabFamily(number=no)
            doc.families.append(current)
        current.words.append(
            VocabWord(
                word=word,
                word_type=str(row.get("Type", "") or "").strip(),
                pronunciation=str(row.get("Pronunciation", "") or "").strip(),
                meaning=str(row.get("Meaning", "") or "").strip(),
            )
        )
    return doc


def csv_to_rows(text: str) -> list[dict]:
    reader = csv.reader(io.StringIO((text or "").strip()))
    out: list[dict] = []
    for i, row in enumerate(reader):
        if not row or not "".join(row).strip():
            continue
        while len(row) < 5:
            row.append("")
        row = row[:5]
        if i == 0 and [c.strip().lower() for c in row[:3]] in (
            ["no.", "word", "type"],
            ["no", "word", "type"],
        ):
            continue
        out.append(dict(zip(["No.", "Word", "Type", "Pronunciation", "Meaning"], row)))
    return out


def _remove_text_runs_preserve_drawings(paragraph) -> None:
    for run in list(paragraph.runs):
        if not run._r.xpath(".//w:drawing | .//w:pict"):
            run._element.getparent().remove(run._element)


def _set_single_run(paragraph, text: str, style: TextStyle, align) -> None:
    _remove_text_runs_preserve_drawings(paragraph)
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_together = False
    fmt.keep_with_next = False
    run = paragraph.add_run(text)
    _apply_run_style(run, style)


def _set_title(doc: Document, data: VocabDocument, theme: Theme) -> None:
    if not doc.tables:
        return
    title = doc.tables[0]
    if len(title.rows) < 2:
        return
    badge = title.rows[0].cells[0]
    badge.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_single_run(badge.paragraphs[0], data.unit_badge, theme.badge, WD_ALIGN_PARAGRAPH.CENTER)

    if len(title.rows[0].cells) >= 3:
        type_cell = title.rows[0].cells[2]
        type_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_single_run(type_cell.paragraphs[0], data.document_type, theme.document_type, WD_ALIGN_PARAGRAPH.LEFT)

    if len(title.rows[1].cells) >= 3:
        main = title.rows[1].cells[2]
        main.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        while len(main.paragraphs) > 1:
            el = main.paragraphs[-1]._element
            el.getparent().remove(el)
        heading = f"{data.section_label} {data.display_heading_unit}: {data.title.upper()}"
        _set_single_run(main.paragraphs[0], heading, theme.title, WD_ALIGN_PARAGRAPH.LEFT)


def _set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        tr_pr.append(node)
    node.set(qn("w:val"), "true")


def _set_cant_split(row, enabled: bool = True) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:cantSplit"))
    if enabled and node is None:
        tr_pr.append(OxmlElement("w:cantSplit"))
    elif not enabled and node is not None:
        tr_pr.remove(node)


def _strip_row_height(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for node in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(node)


def _set_cell_shading(cell, fill: str | None) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if not fill:
        if shd is not None:
            tcpr.remove(shd)
        return
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(fill, "FFFFFF"))


def _set_cell_margins(cell, *, top=55, bottom=55, left=70, right=70) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(int(value)))
        node.set(qn("w:type"), "dxa")


def _set_cell_text(cell, text: str, style: TextStyle, *, center=False) -> None:
    while len(cell.paragraphs) > 1:
        el = cell.paragraphs[-1]._element
        el.getparent().remove(el)
    p = cell.paragraphs[0]
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_together = False
    fmt.keep_with_next = False
    fmt.line_spacing = 1.0
    run = p.add_run(text or "")
    _apply_run_style(run, style)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)


def _set_table_borders(table, color: str) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for name, size in (
        ("top", 8), ("left", 8), ("bottom", 8), ("right", 8),
        ("insideH", 4), ("insideV", 4),
    ):
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), _hex(color, "0F4761"))


def _set_fixed_widths(table, section) -> list[int]:
    content_emu = int(section.page_width - section.left_margin - section.right_margin)
    content_twips = max(1, round(content_emu / EMU_PER_TWIP))

    tblpr = table._tbl.tblPr
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(content_twips))

    ratios = [0.075, 0.235, 0.155, 0.215, 0.320]
    widths = [round(content_twips * r) for r in ratios]
    widths[-1] += content_twips - sum(widths)

    grid_cols = table._tbl.tblGrid.gridCol_lst
    for i, width in enumerate(widths):
        if i < len(grid_cols):
            grid_cols[i].set(qn("w:w"), str(width))

    for row in table.rows:
        for i, cell in enumerate(row.cells[:5]):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:type"), "dxa")
            tcw.set(qn("w:w"), str(widths[i]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    return widths


def _clear_vmerge(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    vm = tcpr.find(qn("w:vMerge"))
    if vm is not None:
        tcpr.remove(vm)


def _populate_table(doc: Document, data: VocabDocument, theme: Theme) -> None:
    if len(doc.tables) < 2:
        raise RuntimeError("Template cần title table và vocabulary table.")
    table = doc.tables[1]
    if len(table.rows) < 3:
        raise RuntimeError("Vocabulary table cần header và hai prototype rows.")

    header = table.rows[0]
    root_proto = table.rows[1]
    child_proto = table.rows[2]

    for row in list(table.rows[3:]):
        table._tbl.remove(row._tr)

    generated = []
    for family in data.families:
        for idx, item in enumerate(family.words):
            proto = root_proto if idx == 0 else child_proto
            table._tbl.append(deepcopy(proto._tr))
            row = table.rows[-1]
            generated.append((row, family.number if idx == 0 else "", item))

    table._tbl.remove(root_proto._tr)
    table._tbl.remove(child_proto._tr)

    _set_fixed_widths(table, doc.sections[0])
    _set_table_borders(table, theme.navy)
    _set_repeat_header(header)
    _set_cant_split(header, True)
    _strip_row_height(header)

    labels = ["No.", "Word", "Type", "Pronunciation", "Meaning"]
    for cell, label in zip(header.cells[:5], labels):
        _clear_vmerge(cell)
        _set_cell_shading(cell, theme.navy)
        _set_cell_text(cell, label, theme.header, center=True)

    for row, number, item in generated:
        _set_cant_split(row, True)
        _strip_row_height(row)
        for cell in row.cells[:5]:
            _clear_vmerge(cell)
        _set_cell_shading(row.cells[0], theme.accent)
        for cell in row.cells[1:5]:
            _set_cell_shading(cell, None)

        _set_cell_text(row.cells[0], number, theme.no_text, center=True)
        _set_cell_text(row.cells[1], item.word, theme.body)
        _set_cell_text(row.cells[2], item.word_type, theme.body)
        _set_cell_text(row.cells[3], item.pronunciation, theme.body)
        _set_cell_text(row.cells[4], item.meaning, theme.body)


def _clear_footer(footer) -> None:
    for child in list(footer._element):
        footer._element.remove(child)


def _no_borders(table) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "nil")


def _inline_page_badge(paragraph, theme: Theme, size_pt: float = 24.0) -> None:
    size_emu = int(round(size_pt * 12700))
    base = _hex(theme.accent, "BF4E14")
    style = theme.page
    hp = max(12, int(round(style.size * 2)))
    bold_xml = "<w:b/><w:bCs/>" if style.bold else ""
    italic_xml = "<w:i/><w:iCs/>" if style.italic else ""
    rpr = (
        f'<w:rPr><w:rFonts w:ascii="{style.font}" w:hAnsi="{style.font}" '
        f'w:eastAsia="{style.font}" w:cs="{style.font}"/>{bold_xml}{italic_xml}'
        f'<w:color w:val="{_hex(style.color, "FFFFFF")}"/>'
        f'<w:sz w:val="{hp}"/><w:szCs w:val="{hp}"/></w:rPr>'
    )
    field_runs = f"""
      <w:r>{rpr}<w:fldChar w:fldCharType="begin"/></w:r>
      <w:r>{rpr}<w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>
      <w:r>{rpr}<w:fldChar w:fldCharType="separate"/></w:r>
      <w:r>{rpr}<w:t>1</w:t></w:r>
      <w:r>{rpr}<w:fldChar w:fldCharType="end"/></w:r>"""

    xml = f"""
    <w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
      <wp:inline distT="0" distB="0" distL="0" distR="0">
        <wp:extent cx="{size_emu}" cy="{size_emu}"/>
        <wp:effectExtent l="0" t="0" r="0" b="0"/>
        <wp:docPr id="910051" name="LX V5 Page Badge"/>
        <wp:cNvGraphicFramePr/>
        <a:graphic>
          <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
            <wps:wsp>
              <wps:cNvSpPr/>
              <wps:spPr>
                <a:xfrm><a:off x="0" y="0"/><a:ext cx="{size_emu}" cy="{size_emu}"/></a:xfrm>
                <a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom>
                <a:gradFill flip="none" rotWithShape="1">
                  <a:gsLst>
                    <a:gs pos="0"><a:srgbClr val="{base}"/></a:gs>
                    <a:gs pos="100000"><a:srgbClr val="{base}"><a:lumMod val="70000"/></a:srgbClr></a:gs>
                  </a:gsLst>
                  <a:lin ang="2700000" scaled="1"/>
                </a:gradFill>
                <a:ln><a:noFill/></a:ln>
              </wps:spPr>
              <wps:txbx>
                <w:txbxContent>
                  <w:p>
                    <w:pPr><w:spacing w:before="0" w:after="0"/><w:jc w:val="center"/></w:pPr>
                    {field_runs}
                  </w:p>
                </w:txbxContent>
              </wps:txbx>
              <wps:bodyPr rot="0" vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0"
                  anchor="ctr" anchorCtr="1">
                <a:prstTxWarp prst="textNoShape"><a:avLst/></a:prstTxWarp>
                <a:noAutofit/>
              </wps:bodyPr>
            </wps:wsp>
          </a:graphicData>
        </a:graphic>
      </wp:inline>
    </w:drawing>"""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run._r.append(etree.fromstring(xml.encode("utf-8")))


def _set_footer(doc: Document, theme: Theme) -> None:
    section = doc.sections[0]
    footer = section.footer
    _clear_footer(footer)

    content_emu = int(section.page_width - section.left_margin - section.right_margin)
    center_emu = int(Inches(0.52))
    side_emu = max(1, (content_emu - center_emu) // 2)

    table = footer.add_table(rows=1, cols=3, width=content_emu)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _no_borders(table)
    widths = [side_emu, center_emu, side_emu]
    for i, width in enumerate(widths):
        table.columns[i].width = width
        table.cell(0, i).width = width
        table.cell(0, i).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    left = table.cell(0, 0)
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("From Learners to Explorers")
    _apply_run_style(r, theme.footer)

    center = table.cell(0, 1)
    _inline_page_badge(center.paragraphs[0], theme)

    right = table.cell(0, 2)
    p1 = right.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r = p1.add_run(f"☎ {theme.phone}")
    _apply_run_style(r, theme.footer)
    p2 = right.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run(f"✉ {theme.email}")
    _apply_run_style(r, theme.footer)

    tail = footer.add_paragraph()
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(0)
    tail.paragraph_format.line_spacing = Pt(1)


def render_docx(
    data: VocabDocument,
    *,
    template_path: str | Path,
    theme: Theme | None = None,
) -> bytes:
    theme = theme or Theme()
    doc = Document(str(template_path))
    _set_title(doc, data, theme)
    _populate_table(doc, data, theme)
    _set_footer(doc, theme)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def inspect_layout(docx_bytes: bytes) -> list[str]:
    problems: list[str] = []
    with ZipFile(io.BytesIO(docx_bytes)) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        footer_xml = "\n".join(
            zf.read(name).decode("utf-8", errors="ignore")
            for name in zf.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )

    if 'w:type="fixed"' not in document_xml:
        problems.append("Table chưa ở fixed layout.")

    doc_root = etree.fromstring(document_xml.encode("utf-8"))
    enabled_keep_next = []
    for node in doc_root.xpath(".//w:keepNext", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}):
        value = node.get(qn("w:val"))
        if value not in ("0", "false", "off"):
            enabled_keep_next.append(node)
    if enabled_keep_next:
        problems.append("Vẫn còn keepNext đang bật trong document.xml.")

    if "PAGE" not in footer_xml:
        problems.append("Footer chưa có PAGE field.")
    if "wp:inline" not in footer_xml:
        problems.append("Page badge chưa dùng inline drawing.")
    if "wp:anchor" in footer_xml and "LX V5 Page Badge" in footer_xml:
        problems.append("Page badge V5 vẫn còn floating anchor.")
    return problems
